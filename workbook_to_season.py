#!/usr/bin/env python3
"""
One-time migration: the Google Sheets schedule engine -> canonical season JSON.

Run:  python3 tools/workbook_to_season.py <workbook.xlsx> <league-id> <out.json>

Writes into data/, which is gitignored. Never writes player data.
"""
import sys, json, datetime
from openpyxl import load_workbook

def iso(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    return None

def hhmm(v):
    if isinstance(v, datetime.time): return v.strftime('%H:%M')
    if isinstance(v, datetime.datetime): return v.strftime('%H:%M')
    return None

def rink(v):
    s = str(v or '').strip()
    for suf in (' Rink', ' Arena', ' Sheet'):
        if s.endswith(suf): s = s[: -len(suf)]
    return s.strip()

def slot_id(d, t, r):
    return f"slot-{d}-{t.replace(':','')}-{''.join(ch for ch in r if ch.isalnum())}"

def main(path, league_id, out):
    wf = load_workbook(path)
    wv = load_workbook(path, data_only=True)
    G = 'Game Schedule' if 'Game Schedule' in wf.sheetnames else 'Schedule'
    P = 'Practices Schedule' if 'Practices Schedule' in wf.sheetnames else 'Practices'

    st = wv['Setup']
    season = {
        'schemaVersion': 1, 'id': 'msu-2026-27', 'name': str(st['B4'].value),
        'leagueId': league_id,
        'startDate': iso(st['B5'].value), 'endDate': iso(st['B6'].value),
        'venue': str(st['B8'].value or ''), 'city': str(st['B9'].value or ''),
        'divisions': [], 'teams': [], 'slots': [], 'games': [], 'practices': [], 'people': [],
    }
    for r in range(15, 19):
        name = st.cell(r, 1).value
        if not name: continue
        gt, pt = st.cell(r, 3).value, st.cell(r, 4).value
        season['divisions'].append({
            'id': str(name), 'name': str(name),
            'gamesPerTeam': int(gt) if gt is not None else None,
            'practicesPerTeam': int(pt) if pt is not None else None,
            'playoffTeams': int(st['B7'].value or 0),
        })
    tm = wv['Teams']
    for r in range(3, 19):
        n, d = tm.cell(r, 2).value, tm.cell(r, 3).value
        if not n: continue
        season['teams'].append({
            'id': str(n), 'divisionId': str(d), 'name': str(n),
            'coach': str(tm.cell(r, 5).value or ''), 'blackoutDates': [], 'allowedWeekdays': [],
        })

    ice = wv['Ice']
    byslot = {}
    for r in range(4, ice.max_row + 1):
        d, t = iso(ice.cell(r, 2).value), hhmm(ice.cell(r, 4).value)
        if not d or not t: continue
        rk = rink(ice.cell(r, 6).value) or 'Unknown'
        use = str(ice.cell(r, 7).value or 'Game').strip().lower()
        use = {'game': 'game', 'practice': 'practice', 'unused': 'unused'}.get(use, 'game')
        sid = slot_id(d, t, rk)
        if sid in byslot: continue
        byslot[sid] = {
            'id': sid, 'date': d, 'start': t,
            'minutes': int(ice.cell(r, 5).value or 60), 'rink': rk, 'use': use,
            'capacity': int(ice.cell(r, 9).value or 1), 'practiceTeams': 1,
            'ownerDivisionId': (str(ice.cell(r, 8).value) if ice.cell(r, 8).value else None),
            'tag': str(ice.cell(r, 11).value or '') if use != 'game' else '',
        }

    gf, gv = wf[G], wv[G]
    TYPE = {'regular': 'regular', 'semifinal': 'semifinal', 'final': 'final'}
    teamnames = {t['id'] for t in season['teams']}
    n = 0
    for r in range(4, gf.max_row + 1):
        d, t = iso(gv.cell(r, 2).value), hhmm(gv.cell(r, 3).value)
        div = gv.cell(r, 4).value
        typ = str(gv.cell(r, 11).value or '').strip().lower()
        if typ not in TYPE: continue
        home, away = gv.cell(r, 6).value, gv.cell(r, 7).value
        if not d or not t or not div or (not home and not away): continue
        rk = rink(gv.cell(r, 5).value) or 'Unknown'
        sid = slot_id(d, t, rk)
        if sid not in byslot:
            print(f'  WARN row {r}: no Ice slot for {d} {t} {rk}', file=sys.stderr)
            continue
        n += 1
        hs, as_ = str(home or ''), str(away or '')
        season['games'].append({
            'id': f'g{n}', 'divisionId': str(div), 'slotId': sid,
            'homeTeamId': hs if hs in teamnames else None,
            'awayTeamId': as_ if as_ in teamnames else None,
            'homeLabel': '' if hs in teamnames else hs,
            'awayLabel': '' if as_ in teamnames else as_,
            'type': TYPE[typ],
            'homeGoals': gv.cell(r, 8).value, 'awayGoals': gv.cell(r, 9).value,
        })

    pv = wv[P]
    pract = {}
    for r in range(4, pv.max_row + 1):
        d, t = iso(pv.cell(r, 2).value), hhmm(pv.cell(r, 3).value)
        team, div = pv.cell(r, 7).value, pv.cell(r, 6).value
        if not d or not t or not team: continue
        rk = rink(pv.cell(r, 5).value) or 'Unknown'
        sid = slot_id(d, t, rk)
        if sid not in byslot:
            print(f'  WARN practice row {r}: no Ice slot for {d} {t} {rk}', file=sys.stderr)
            continue
        kind = 'evaluation' if str(pv.cell(r, 1).value or '').strip().upper() == 'EVALS' else 'practice'
        pract.setdefault((sid, str(div), kind), []).append(str(team))
    for i, ((sid, div, kind), teams) in enumerate(sorted(pract.items()), 1):
        season['practices'].append({
            'id': f'pr{i}', 'divisionId': div, 'slotId': sid, 'teamIds': teams, 'kind': kind,
        })
        byslot[sid]['practiceTeams'] = max(byslot[sid]['practiceTeams'], len(teams))

    season['slots'] = sorted(byslot.values(), key=lambda s: (s['date'], s['start'], s['rink']))
    with open(out, 'w') as f: json.dump(season, f, indent=2)
    print(f"{out}: {len(season['slots'])} slots, {len(season['games'])} games, "
          f"{len(season['practices'])} practice groups, {len(season['teams'])} teams")

if __name__ == '__main__':
    main(*sys.argv[1:4])
